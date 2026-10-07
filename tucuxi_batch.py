# -*- coding: utf-8 -*-
"""
tucuxi_batch.py — Tính Bayes HÀNG LOẠT từ file Excel (theo mẫu 'data dữ liệu') bằng engine tucuxi_engine.

Cho mỗi bệnh nhân, mỗi lần TDM N, mỗi mô hình (tucuxi.goti / .collin / .thomson / .yamamoto):
  • tiên nghiệm tại T(N) (η = 0)                       — lần TDM 1 dùng làm "dự đoán" quần thể
  • dự đoán Bayes N−1 → N                              — N ≥ 2: khớp bằng mẫu N−1 (hoặc 1..N−1 nếu cumulative=True),
                                                         liều thực tế đến T(N), SCr chỉ đến T(N−1) (không dùng SCr "tương lai")
  • độ khớp: nồng độ sau Bayes tại T(N) khi khớp bằng mẫu N (hoặc 1..N)
Chỉ số: rBias% = (Cdđ − Cđo)·2/(Cdđ + Cđo)·100, APE%, rRMSE%, MdAPE, P20/P30.
Quy ước đọc Excel giữ nguyên như bộ dữ liệu mẫu: mẫu thiếu SCr → dùng SCr kế trước; cột "Đến" là mốc LOẠI TRỪ nếu
dòng liều kế tiếp bắt đầu đúng mốc đó; "Ngưng" = 1 liều.
"""
from __future__ import annotations

import datetime
import math
import statistics
import traceback
from typing import Callable, Dict, List, Optional

import openpyxl

import tucuxi_engine as te

TD = datetime.timedelta
DEFAULT_MODELS = ["tucuxi.goti", "tucuxi.collin", "tucuxi.thomson", "tucuxi.yamamoto"]


# ---------------------------------------------------------------------------------------------
def read_cases(path: str) -> List[dict]:
    """Đọc sheet đầu tiên 'data dữ liệu' (dữ liệu từ hàng 3). Tuổi: nam cột B, nữ cột C; D cân nặng; E chiều cao;
    F thời điểm / từ; G đến; H liều mg; I khoảng đưa liều (h); J thời gian truyền (h); K SCr µmol/L; M nồng độ mg/L."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["data dữ liệu"] if "data dữ liệu" in wb.sheetnames else wb[wb.sheetnames[0]]
    cases: List[dict] = []
    cur = None
    for r in range(3, ws.max_row + 1):
        v = [ws.cell(r, c).value for c in range(1, 15)]
        if v[0] is not None:
            male_age, female_age = v[1], v[2]
            try:
                age = float(male_age if male_age is not None else female_age)
                weight = float(v[3])
            except (TypeError, ValueError):
                cur = None
                cases.append(dict(msyt=v[0], skip="Thiếu tuổi/cân nặng"))
                continue
            cur = dict(msyt=v[0], row=r, male=male_age is not None, age=age, weight=weight,
                       height=v[4], dose_rows=[], scr=[], samples=[], warnings=[])
            cases.append(cur)
            if isinstance(v[5], datetime.datetime) and isinstance(v[10], (int, float)):
                cur["scr"].append((v[5], float(v[10])))
            continue
        if cur is None or not isinstance(v[5], datetime.datetime):
            continue
        F, G, mg, tau, tinf, scr, cobs = v[5], v[6], v[7], v[8], v[9], v[10], v[12]
        if isinstance(mg, (int, float)):
            mg = float(mg) * 1000.0 if mg < 20 else float(mg)             # < 20 → gam
            if not isinstance(tinf, (int, float)):
                tinf = 2.0
                cur["warnings"].append(f"Hàng {r}: thiếu thời gian truyền, giả định 2 giờ")
            cur["dose_rows"].append(dict(F=F, G=G, mg=mg, tau=float(tau), tinf=float(tinf), row=r))
        else:
            if isinstance(scr, (int, float)):
                cur["scr"].append((F, float(scr)))
            if isinstance(cobs, (int, float)):
                if not isinstance(scr, (int, float)):
                    prev = [s for s in cur["scr"] if s[0] <= F]
                    if prev:
                        cur["scr"].append((F, prev[-1][1]))
                        cur["warnings"].append(f"Hàng {r}: thiếu SCr, dùng SCr kế trước")
                cur["samples"].append((F, float(cobs), r))
    for c in cases:
        if "skip" in c:
            continue
        c["scr"].sort(key=lambda x: x[0])
        c["samples"].sort(key=lambda x: x[0])
        c["doses"] = _expand_doses(c["dose_rows"])
        if not c["samples"]:
            c["skip"] = "Chưa có nồng độ"
        elif not c["doses"]:
            c["skip"] = "Chưa có liều"
        elif not c["scr"]:
            c["skip"] = "Chưa có SCr"
        elif c["weight"] > 120 and not c["height"]:
            c["skip"] = "Cân nặng > 120 kg và chiều cao trống (nghi nhập nhầm cột)"
    return cases


def _expand_doses(dose_rows):
    out = []
    for i, d in enumerate(dose_rows):
        nxt = dose_rows[i + 1]["F"] if i + 1 < len(dose_rows) else None
        if not isinstance(d["G"], datetime.datetime):
            out.append((d["F"], d["mg"], d["tinf"], d["tau"]))
            continue
        t = d["F"]
        while t <= d["G"]:
            if t < d["G"] or nxt is None or nxt > d["G"]:
                out.append((t, d["mg"], d["tinf"], d["tau"]))
            t += TD(hours=d["tau"])
    out.sort(key=lambda x: x[0])
    return out


# ---------------------------------------------------------------------------------------------
def rbias(c_pred: float, c_obs: float) -> float:
    return (c_pred - c_obs) * 2.0 / (c_pred + c_obs) * 100.0


def _stats(vals: List[float], apes: List[float]) -> dict:
    n = len(vals)
    if n == 0:
        return dict(n=0)
    return dict(n=n, rbias=sum(vals) / n, rrmse=math.sqrt(sum(x * x for x in vals) / n),
                mdape=statistics.median(apes), p20=100.0 * sum(a <= 20 for a in apes) / n,
                p30=100.0 * sum(a <= 30 for a in apes) / n)


def run_case_model(case: dict, model_key: str, cumulative: bool = False, gram_positive: bool = True,
                   drug_dir: Optional[str] = None) -> List[dict]:
    pat = te.TuPatient(case["age"], case["male"], case["weight"], gram_positive=gram_positive)
    S = case["samples"]
    doses3 = [(t, mg, ti) for t, mg, ti, _ in case["doses"]]
    out = []
    for n, (t, cobs, row) in enumerate(S, start=1):
        rec = dict(msyt=case["msyt"], n=n, row=row, t=t, cobs=cobs, model=model_key)
        ds = [d for d in doses3 if d[0] <= t]
        if not ds:
            rec["warn"] = "Chưa có liều trước thời điểm đo"
            out.append(rec)
            continue
        scr_n = [(a, b) for a, b in case["scr"] if a <= t]
        smp_n = [(s[0], s[1]) for s in S[:n]] if cumulative else [(t, cobs)]
        try:
            fit = te.tucuxi_fit(model_key, pat, ds, scr_n, smp_n, drug_dir=drug_dir)
            rec["c_prior"] = fit.predict([t], kind="prior")[0]
            rec["c_fit"] = fit.predict([t], kind="post")[0]
            rec.update(CL=fit.CL_optimized, V1=fit.Vc_optimized, V2=fit.Vp_optimized, Q=fit.Q_value,
                       CL_prior=fit.CL_prior, nll=fit.nll)
            last = [d for d in case["doses"] if d[0] <= t][-1]
            rec["dose_mg"], rec["tau_h"] = last[1], last[3]
            rec["auc24"] = last[1] * 24.0 / (last[3] * fit.CL_optimized)
            if n == 1:
                rec["c_pred"] = rec["c_prior"]
                rec["pred_kind"] = "Tiên nghiệm"
            else:
                tp = S[n - 2][0]
                smp_p = [(s[0], s[1]) for s in S[:n - 1]] if cumulative else [(tp, S[n - 2][1])]
                scr_p = [(a, b) for a, b in case["scr"] if a <= tp]
                fp = te.tucuxi_fit(model_key, pat, ds, scr_p, smp_p, drug_dir=drug_dir)
                rec["c_pred"] = fp.predict([t])[0]
                rec["pred_kind"] = "Bayes N−1→N"
            rec["rbias_pred"] = rbias(rec["c_pred"], cobs)
            rec["ape_pred"] = abs(rec["c_pred"] - cobs) / cobs * 100.0
            rec["rbias_fit"] = rbias(rec["c_fit"], cobs)
        except Exception as e:                                   # 1 ca lỗi không làm hỏng cả lô
            rec["warn"] = f"Lỗi tính: {e}"
        out.append(rec)
    return out


def summarize(records: List[dict]) -> Dict[str, dict]:
    res = {}
    for key, flt in (("Tiên nghiệm (TDM 1)", lambda r: r["n"] == 1),
                     ("Bayes N−1→N (TDM ≥ 2)", lambda r: r["n"] >= 2)):
        rs = [r for r in records if flt(r) and "rbias_pred" in r]
        res[key] = _stats([r["rbias_pred"] for r in rs], [r["ape_pred"] for r in rs])
    rs = [r for r in records if "rbias_fit" in r]
    res["Độ khớp sau Bayes (trong mẫu)"] = _stats(
        [r["rbias_fit"] for r in rs], [abs(r["c_fit"] - r["cobs"]) / r["cobs"] * 100.0 for r in rs])
    return res


def run_batch(xlsx_in: str, xlsx_out: str, models: Optional[List[str]] = None, cumulative: bool = False,
              gram_positive: bool = True, progress: Optional[Callable[[int, int, str], None]] = None) -> dict:
    models = models or DEFAULT_MODELS
    cases = read_cases(xlsx_in)
    ok_cases = [c for c in cases if "skip" not in c]
    skipped = [c for c in cases if "skip" in c]
    total = len(ok_cases) * len(models)
    done = 0
    all_rec: Dict[str, List[dict]] = {m: [] for m in models}
    for c in ok_cases:
        for m in models:
            all_rec[m] += run_case_model(c, m, cumulative, gram_positive)
            done += 1
            if progress:
                progress(done, total, f"{c['msyt']} — {m}")
    _write_excel(xlsx_out, models, all_rec, cases, skipped, cumulative)
    return dict(n_cases=len(ok_cases), n_skipped=len(skipped), models=models, out=xlsx_out,
                summary={m: summarize(all_rec[m]) for m in models})


def _write_excel(path, models, all_rec, cases, skipped, cumulative):
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kết quả TDM"
    head = ["MSYT", "Lần TDM", "Hàng Excel", "Thời điểm đo", "C đo (mg/L)", "Mô hình", "Loại dự đoán",
            "C dự đoán (mg/L)", "rBias %", "APE %", "C tiên nghiệm", "C sau Bayes (khớp)", "rBias khớp %",
            "CL (L/h)", "V1 (L)", "V2 (L)", "Q (L/h)", "CL tiên nghiệm", "Liều đang dùng (mg)", "τ (h)",
            "AUC24 (liều đang dùng)", "Ghi chú"]
    ws.append(head)
    for m in models:
        for r in all_rec[m]:
            ws.append([r["msyt"], r["n"], r["row"], r["t"], r["cobs"], m, r.get("pred_kind"), r.get("c_pred"),
                       r.get("rbias_pred"), r.get("ape_pred"), r.get("c_prior"), r.get("c_fit"), r.get("rbias_fit"),
                       r.get("CL"), r.get("V1"), r.get("V2"), r.get("Q"), r.get("CL_prior"), r.get("dose_mg"),
                       r.get("tau_h"), r.get("auc24"), r.get("warn")])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        row[3].number_format = "yyyy-mm-dd hh:mm"
        for i in (7, 10, 11, 13, 14, 15, 16, 17, 20):
            row[i].number_format = "0.00"
        for i in (8, 9, 12):
            row[i].number_format = "0.0"
    ws.freeze_panes = "A2"
    for i, w in enumerate([12, 8, 9, 17, 10, 16, 14, 12, 9, 9, 12, 12, 11, 9, 9, 9, 9, 11, 11, 7, 12, 40], start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    ss = wb.create_sheet("Thống kê")
    ss.append(["Mô hình", "Nhóm so sánh", "n", "rBias % (TB)", "rRMSE %", "MdAPE %", "P20 %", "P30 %"])
    for m in models:
        for grp, s in summarize(all_rec[m]).items():
            if s.get("n"):
                ss.append([m, grp, s["n"], s["rbias"], s["rrmse"], s["mdape"], s["p20"], s["p30"]])
            else:
                ss.append([m, grp, 0])
    for c in ss[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for row in ss.iter_rows(min_row=2):
        for i in range(3, 8):
            row[i].number_format = "0.0"
    for i, w in enumerate([18, 32, 6, 13, 10, 10, 8, 8], start=1):
        ss.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    ws3 = wb.create_sheet("Cảnh báo")
    ws3.append(["MSYT", "Nội dung"])
    for c in skipped:
        ws3.append([c["msyt"], "Bỏ qua: " + c["skip"]])
    for c in cases:
        for w in c.get("warnings", []):
            ws3.append([c["msyt"], w])
    ws3.column_dimensions["B"].width = 80

    ws4 = wb.create_sheet("Ghi chú phương pháp")
    notes = [
        "Engine Python mô phỏng Tucuxi-core, đọc nguyên file .tdd (không sửa mô hình).",
        "Một bộ η cho toàn bộ lịch sử; CL cập nhật theo ngày theo SCr nội suy tuyến tính; clcr = Cockcroft–Gault của Tucuxi.",
        "Chế độ khớp: " + ("CỘNG DỒN (mọi nồng độ đến T(N))" if cumulative else "MỘT nồng độ (chỉ mẫu N) — như bộ dữ liệu mẫu"),
        "Lần TDM 1: dự đoán = tiên nghiệm quần thể. Lần ≥ 2: dự đoán từ thông số hậu nghiệm lần N−1, dùng liều thực tế, SCr đến T(N−1).",
        "rBias% = (Cdđ − Cđo)·2/(Cdđ + Cđo)·100;  rRMSE% = √(TB rBias²).",
        "Đây là công cụ nghiên cứu/đối chiếu — chưa thay thế phê duyệt lâm sàng.",
    ]
    for n in notes:
        ws4.append([n])
    ws4.column_dimensions["A"].width = 130
    wb.save(path)
